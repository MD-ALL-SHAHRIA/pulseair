"""Lay the thesis out in the official AIUB thesis template.

``build_thesis`` decides *what* the thesis says; this module decides only *how it looks*.
The template file is copied and filled in rather than imitated, so its logo, styles,
section breaks, footers ("CSC 4298 Thesis/Project | Department of Computer Science |
page"), page-number formats and Declaration/Approval layouts are the template's own.

Values below were read from the template's ``word/styles.xml``, ``numbering.xml`` and
``document.xml`` (A4 11910x16840 twips, 1440-twip margins, Times New Roman Normal
style), and the visible layout is reproduced as the template pages show it: a chapter
opens with "Chapter N" and its title, both 20 pt bold Heading 1, over a horizontal rule;
sections are 14 pt bold; body text is 12 pt Times New Roman, justified, 1.5 lines,
indented 0.25 in. Captions carry SEQ fields so Word's List of Figures / List of Tables
fields can be refreshed (Ctrl+A, F9); Word is also told to refresh them on open.
"""
from __future__ import annotations

import copy
import itertools
import re
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Twips
from docx.table import Table

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "docs" / "template" / "Thesis Report Template for Spring 2025-26 2.docx"
TEXT_W = 11910 - 2 * 1440          # usable width in twips
BULLET_NUM_ID = "64"              # the template's own bullet list (Times New Roman "o")
FONT = "Times New Roman"
_ids = itertools.count(9000)

W = lambda t: qn("w:" + t)        # noqa: E731

PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "widowControl", "numPr",
             "pBdr", "tabs", "suppressAutoHyphens", "autoSpaceDE", "autoSpaceDN", "spacing",
             "ind", "contextualSpacing", "jc", "textAlignment", "outlineLvl", "rPr", "sectPr"]
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "noProof", "color", "sz",
             "szCs", "u", "vertAlign"]


# ------------------------------------------------------------------ xml helpers
def el(tag, **attrs):
    e = OxmlElement(tag if ":" in tag else "w:" + tag)
    for k, v in attrs.items():
        e.set(W(k), str(v))
    return e


def _sort(parent, order):
    kids = sorted(list(parent), key=lambda x: order.index(x.tag.split("}")[1])
                  if x.tag.split("}")[1] in order else len(order))
    for k in kids:
        parent.remove(k)
    for k in kids:
        parent.append(k)


def ppr(p):
    x = p.find(W("pPr"))
    if x is None:
        x = el("pPr")
        p.insert(0, x)
    return x


def ppr_set(p, child):
    pr = ppr(p)
    for x in pr.findall(child.tag):
        pr.remove(x)
    pr.append(child)
    _sort(pr, PPR_ORDER)


def run_fmt(r, size=12, bold=False, italic=False, color="000000"):
    """Times New Roman, explicit size, black unless told otherwise."""
    rpr = r.find(W("rPr"))
    if rpr is None:
        rpr = el("rPr")
        r.insert(0, rpr)
    for t in ("rFonts", "b", "bCs", "i", "iCs", "sz", "szCs"):
        for x in rpr.findall(W(t)):
            rpr.remove(x)
    rpr.append(el("rFonts", ascii=FONT, hAnsi=FONT, cs=FONT, eastAsia=FONT))
    if bold:
        rpr.append(el("b")); rpr.append(el("bCs"))
    if italic:
        rpr.append(el("i")); rpr.append(el("iCs"))
    if color and rpr.find(W("color")) is None:
        rpr.append(el("color", val=color))
    half = int(round(size * 2))
    rpr.append(el("sz", val=half)); rpr.append(el("szCs", val=half))
    _sort(rpr, RPR_ORDER)


def make_run(text, **fmt):
    r = el("r")
    t = el("t"); t.text = text; t.set(qn("xml:space"), "preserve")
    r.append(t)
    run_fmt(r, **fmt)
    return r


def ptext(p):
    return "".join(t.text or "" for t in p.iter(W("t")))


def clear_runs(p):
    for x in list(p):
        if x.tag != W("pPr"):
            p.remove(x)


def field(instr, cached, **fmt):
    out = []
    for kind in ("begin", None, "separate", "result", "end"):
        if kind == "result":
            out.append(make_run(cached, **fmt)); continue
        r = el("r")
        if kind is None:
            t = el("instrText"); t.text = f" {instr} "; t.set(qn("xml:space"), "preserve"); r.append(t)
        else:
            r.append(el("fldChar", fldCharType=kind))
        run_fmt(r, **fmt)
        out.append(r)
    return out


# ------------------------------------------------------------------ paragraph formats
def format_body(p, *, align=None, small=False):
    """Template body text: 12 pt, justified, 1.5 lines, 0.25 in left indent."""
    ppr_set(p._p, el("pStyle", val="BodyText"))
    if small:   # source notes, reference entries: single-spaced, unindented
        ppr_set(p._p, el("spacing", before=0, after=80, line=264, lineRule="auto"))
    else:
        ppr_set(p._p, el("spacing", before=0, after=120, line=360, lineRule="auto"))
        ppr_set(p._p, el("ind", left=360))
    ppr_set(p._p, el("jc", val={"center": "center", "right": "right"}.get(align, "both")))


def bullet(p):
    ppr_set(p._p, el("pStyle", val="ListParagraph"))
    num = el("numPr"); num.append(el("ilvl", val=0)); num.append(el("numId", val=BULLET_NUM_ID))
    ppr_set(p._p, num)
    ppr_set(p._p, el("spacing", before=0, after=60, line=360, lineRule="auto"))
    ppr_set(p._p, el("ind", left=720, hanging=360))
    ppr_set(p._p, el("jc", val="both"))


def caption(p, kind, n, label_rest):
    """'Figure n: …' / 'Table n: …' with a SEQ field, Caption style (black TNR)."""
    ppr_set(p._p, el("pStyle", val="Caption"))
    ppr_set(p._p, el("jc", val="center"))
    ppr_set(p._p, el("spacing", before=60 if kind == "Figure" else 240,
                     after=0 if kind == "Figure" else 80, line=240, lineRule="auto"))
    if kind == "Table":
        ppr_set(p._p, el("keepNext"))
    p._p.append(make_run(f"{kind} ", size=11, bold=True))
    for r in field(f"SEQ {kind} \\* ARABIC", str(n), size=11, bold=True):
        p._p.append(r)
    p._p.append(make_run(": ", size=11, bold=True))


def style_table(t, widths_in=None, header=True):
    tbl = t._tbl
    t.style = "Table Grid"
    pr = tbl.tblPr
    for tag in ("jc", "tblLayout", "tblW"):
        for x in pr.findall(W(tag)):
            pr.remove(x)
    ncol = len(t.columns)
    if widths_in:
        tw = [int(w * 1440) for w in widths_in]
        scale = min(1.0, (TEXT_W - 200) / sum(tw))
        tw = [int(w * scale) for w in tw]
    else:
        tw = [int((TEXT_W - 200) / ncol)] * ncol
    pr.append(el("tblW", w=sum(tw), type="dxa"))
    pr.append(el("jc", val="center"))
    pr.append(el("tblLayout", type="fixed"))
    _sort(pr, ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
               "shd", "tblLayout", "tblCellMar", "tblLook"])
    grid = tbl.find(W("tblGrid"))
    for gc, w in zip(grid.findall(W("gridCol")), tw):
        gc.set(W("w"), str(w))
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(el("cantSplit"))
        if ri == 0 and header:
            trpr.append(el("tblHeader"))
        for ci, cell in enumerate(row.cells):
            cell.width = Twips(tw[ci])
            if ri == 0 and header:
                cell._tc.get_or_add_tcPr().append(el("shd", val="clear", color="auto", fill="D9D9D9"))
            for p in cell.paragraphs:
                ppr_set(p._p, el("pStyle", val="Normal"))
                ppr_set(p._p, el("spacing", before=20, after=20, line=240, lineRule="auto"))
                ppr_set(p._p, el("jc", val="center" if ri == 0 and header else "left"))


# ------------------------------------------------------------------ headings
class Layout:
    """Holds the template prototypes and places headings the way the template does."""

    def __init__(self, doc):
        self.doc = doc
        body = doc.element.body
        paras = [e for e in body if e.tag == W("p")]
        find = lambda s: next(e for e in paras if ptext(e).strip() == s)  # noqa: E731
        self.ch_label = copy.deepcopy(find("Chapter 1"))
        self.ch_title = copy.deepcopy(find("Chapter 1").getnext())
        self.refs_head = copy.deepcopy(find("References"))
        self.app_head = copy.deepcopy(find("Appendix"))
        for proto in (self.ch_title, self.refs_head, self.app_head):
            for b in proto.findall(W("bookmarkStart")) + proto.findall(W("bookmarkEnd")):
                proto.remove(b)

    def _append(self, e):
        self.doc.element.body.find(W("sectPr")).addprevious(e)
        return e

    def chapter(self, number, title):
        lab = copy.deepcopy(self.ch_label)
        ts = list(lab.iter(W("t")))         # the template splits "Chapter 1" across runs
        ts[0].text = f"Chapter {number}"
        ts[0].set(qn("xml:space"), "preserve")
        for t in ts[1:]:
            t.text = ""
        for dp in lab.iter(qn("wp:docPr")):     # drawing ids must stay unique
            dp.set("id", str(next(_ids)))
        ppr_set(lab, el("pageBreakBefore"))
        ppr_set(lab, el("keepNext"))
        ttl = copy.deepcopy(self.ch_title)
        clear_runs(ttl)
        ttl.append(make_run(title, size=20, bold=True))
        self._append(lab)
        return self._append(ttl)

    def unnumbered(self, text, appendix=False):
        h = copy.deepcopy(self.app_head if appendix else self.refs_head)
        clear_runs(h)
        h.append(make_run(text, size=24 if appendix else 20, bold=True))
        ppr_set(h, el("pageBreakBefore"))
        ppr_set(h, el("spacing", before=0, after=240))
        return self._append(h)

    def section(self, text, level):
        """level 2 = '1.1 …' (14 pt), level 3 = '4.15.1 …' (12 pt). The numbers are part of
        the thesis text, so the template's auto-numbering is not applied on top of them."""
        p = el("p")
        ppr_set(p, el("pStyle", val="Heading1"))
        ppr_set(p, el("keepNext"))
        ppr_set(p, el("spacing", before=240 if level == 2 else 180, after=240 if level == 2 else 120,
                      line=360, lineRule="auto"))
        ppr_set(p, el("ind", left=0))
        ppr_set(p, el("outlineLvl", val=level - 1))
        p.append(make_run(text, size=14 if level == 2 else 12, bold=True))
        return self._append(p)


# ------------------------------------------------------------------ front matter
def fill_front_matter(doc, *, title, authors, supervisor, co_supervisor, acknowledgement,
                      abstract_chunks, keywords, abbreviations):
    body = doc.element.body

    def paras():
        return [e for e in body if e.tag == W("p")]

    def find(pred):
        return next(e for e in paras() if pred(ptext(e)))

    def set_text(p, text, size=None, bold=None):
        rs = p.findall(W("r"))
        rp = rs[0].find(W("rPr")) if rs else None
        if size is None and rp is not None and rp.find(W("sz")) is not None:
            size = int(rp.find(W("sz")).get(W("val"))) / 2
        if bold is None:
            bold = rp is not None and rp.find(W("b")) is not None and rp.find(W("b")).get(W("val")) != "0"
        clear_runs(p)
        p.append(make_run(text, size=size or 12, bold=bold))

    # title page ------------------------------------------------------------
    set_text(find(lambda s: s.strip() == "Title of the Thesis"), title, size=22)
    for a, p in zip(authors, [e for e in paras() if "full name" in ptext(e)]):
        set_text(p, f"{a} (xx-xxxxx-x)", size=17)
    for e in body.iter(W("color")):         # red template placeholders become black text
        if e.get(W("val")).upper() in ("FF0000", "EE0000", "C00000"):
            e.set(W("val"), "000000")
    sem = find(lambda s: "2025-2026 Semester" in s)
    set_text(sem, "[Spring/Fall] 2025-2026 Semester", size=14)
    set_text(find(lambda s: s.startswith("Submission Date")), "Submission Date: [Month, Year]", size=14)
    first = doc.sections[0]                  # no footer on the cover
    first.different_first_page_header_footer = True
    first.first_page_footer.is_linked_to_previous = False

    # declaration: template wording verbatim; only the instruction line goes ------
    body.remove(find(lambda s: s.startswith("(All candidates")))
    sig = next(Table(e, doc) for e in body if e.tag == W("tbl") and "Name of Group Member" in e.xpath("string(.)"))
    k = 0
    for row in sig.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                s = ptext(p._p).strip()
                if s.startswith("Name of Group Member"):
                    set_text(p._p, authors[k]); k += 1
                elif s == "AIUB ID":
                    set_text(p._p, "xx-xxxxx-x")
                elif s == "Department":
                    set_text(p._p, "Department of Computer Science")

    # approval ---------------------------------------------------------------
    ap = find(lambda s: s.startswith("The thesis titled"))
    txt = ptext(ap).replace("Thesis title goes here", title).replace("(date of defense)", "[date of defense]")
    pre, rest = txt.split("“", 1)
    ttl, post = rest.split("”", 1)
    clear_runs(ap)
    ap.append(make_run(pre + "“")); ap.append(make_run(ttl, bold=True)); ap.append(make_run("”" + post))
    apt = next(Table(e, doc) for e in body if e.tag == W("tbl") and "Name of the Supervisor" in e.xpath("string(.)"))
    for row in apt.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                s = ptext(p._p).strip()
                if s == "Name of the Supervisor":
                    set_text(p._p, supervisor)
                elif s == "Name of the External":
                    set_text(p._p, co_supervisor)
                elif s.startswith("Rank") and "Supervisor" in s:
                    set_text(p._p, "[Rank] & Supervisor", bold=False)
                elif s.startswith("Rank") and "External" in s:
                    set_text(p._p, "[Rank] & Co-Supervisor", bold=False)
    e = apt._tbl.getnext()   # spacers that would push the page break onto a blank page
    while e is not None and e.tag == W("p") and e.find(".//" + W("br")) is None and not ptext(e).strip():
        nxt = e.getnext(); body.remove(e); e = nxt

    def body_par(text, small=False):
        p = el("p")
        ppr_set(p, el("pStyle", val="BodyText"))
        ppr_set(p, el("spacing", before=0, after=120, line=360, lineRule="auto"))
        ppr_set(p, el("ind", left=360))
        ppr_set(p, el("jc", val="both"))
        p.append(make_run(text))
        return p

    def drop_until_section_break(head, keep_tables=False):
        e = head.getnext()
        while e is not None and not (e.tag == W("p") and (e.find(".//" + W("sectPr")) is not None
                                                          or ptext(e).strip() in ("Abstract", "Keywords"))):
            nxt = e.getnext()
            if e.tag == W("p") and e.find(".//" + W("br")) is None:
                body.remove(e)
            elif e.tag == W("tbl") and not keep_tables:
                body.remove(e)
            e = nxt

    # acknowledgement ---------------------------------------------------------
    ack = find(lambda s: s.strip() == "Acknowledgement")
    e = ack.getnext()
    while ptext(e).strip() != "Table of Content":
        nxt = e.getnext()
        if e.tag == W("p") and e.find(".//" + W("br")) is None:
            body.remove(e)
        e = nxt
    anchor = ack
    for chunk in acknowledgement:
        anchor.addnext(body_par(chunk)); anchor = anchor.getnext()

    # TOC / List of Figures / List of Tables as real fields ---------------------
    sdt = next(e for e in body if e.tag == W("sdt"))
    p = el("p")
    for r in field('TOC \\o "1-3" \\h \\z \\u', "Press Ctrl+A, then F9, to build the Table of Content."):
        p.append(r)
    sdt.addprevious(p)
    body.remove(sdt)
    for name, kind in (("List of Figures", "Figure"), ("List of Tables", "Table")):
        head = find(lambda s, n=name: s.strip() == n)
        drop_until_section_break(head)
        p = el("p")
        for r in field(f'TOC \\h \\z \\c "{kind}"', f"Press Ctrl+A, then F9, to build the List of {kind}s."):
            p.append(r)
        head.addnext(p)

    # abbreviations: the template's own two-column borderless table -------------
    ab_head = find(lambda s: s.strip() == "List of Abbreviations")
    abt = next(Table(e, doc) for e in body if e.tag == W("tbl") and "Human Computer Interaction" in e.xpath("string(.)"))
    proto = copy.deepcopy(abt.rows[0]._tr)
    for r in list(abt.rows):
        abt._tbl.remove(r._tr)
    for abbr, full in abbreviations:
        tr = copy.deepcopy(proto)
        for tc, val in zip(tr.findall(W("tc")), (abbr, full)):
            set_text(tc.find(W("p")), val, size=12)
        abt._tbl.append(tr)
    tbl = abt._tbl
    tbl.getparent().remove(tbl)
    drop_until_section_break(ab_head)
    ab_head.addnext(tbl)

    # abstract + keywords --------------------------------------------------------
    abs_head = find(lambda s: s.strip() == "Abstract")
    kw_head = find(lambda s: s.strip() == "Keywords")
    e = abs_head.getnext()
    while e is not kw_head:
        nxt = e.getnext(); body.remove(e); e = nxt
    anchor = abs_head
    for chunk in abstract_chunks:
        anchor.addnext(body_par(chunk)); anchor = anchor.getnext()
    drop_until_section_break(kw_head)
    kw_head.addnext(body_par(keywords))
    for name in ("Table of Content", "List of Abbreviations", "Abstract"):
        ppr_set(find(lambda s, n=name: s.strip() == n), el("pageBreakBefore"))


def strip_template_body(doc):
    """Remove the template's chapters/references/appendix placeholders, keeping the
    section break that closes the main body (it carries the Arabic page numbering) and
    the final section (the template's appendix section)."""
    body = doc.element.body
    elems = list(body)
    start = next(i for i, e in enumerate(elems) if e.tag == W("p") and ptext(e).strip() == "Keywords")
    start = next(i for i in range(start, len(elems)) if elems[i].find(".//" + W("sectPr")) is not None) + 1
    main_break = next(i for i in range(start, len(elems)) if elems[i].tag == W("p")
                      and elems[i].find(".//" + W("sectPr")) is not None)
    for e in elems[start:main_break] + elems[main_break + 1:-1]:
        body.remove(e)
    brk = elems[main_break]
    sp = brk.find(".//" + W("sectPr"))
    for x in sp.findall(W("pgNumType")):
        sp.remove(x)
    sp.find(W("cols")).addprevious(el("pgNumType", fmt="decimal", start=1))
    return brk


def finish(doc, main_break, n_before_body, n_before_appendix):
    """Move the main body (chapters → references) in front of the main-body section
    break; the appendices stay after it, in the template's appendix section. Elements
    stay in the same package, so every image part keeps its relationship."""
    body = doc.element.body
    elems = list(body)[:-1]
    main = elems[n_before_body:n_before_appendix]
    for e in main:
        main_break.addprevious(e)
    styles = doc.styles.element
    for sid, size, bold in (("Caption", 11, False), ("TOC1", 12, True), ("TOC2", 12, False),
                            ("TOC3", 12, False), ("TableofFigures", 12, False)):
        st = styles.find(f"{W('style')}[@{W('styleId')}='{sid}']")
        if st is None:
            continue
        rpr = st.find(W("rPr"))
        if rpr is None:
            rpr = el("rPr"); st.append(rpr)
        for x in list(rpr):
            rpr.remove(x)
        rpr.append(el("rFonts", ascii=FONT, hAnsi=FONT, cs=FONT, eastAsia=FONT))
        if bold:
            rpr.append(el("b"))
        rpr.append(el("color", val="000000"))
        rpr.append(el("sz", val=size * 2)); rpr.append(el("szCs", val=size * 2))
    settings = doc.settings.element
    if settings.find(W("updateFields")) is None:
        settings.append(el("updateFields", val="true"))


def add_hyperlink(doc, p, url, size):
    rid = doc.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = el("hyperlink"); h.set(qn("r:id"), rid)
    r = make_run(url, size=size, color="0563C1")
    r.find(W("rPr")).append(el("u", val="single"))
    _sort(r.find(W("rPr")), RPR_ORDER)
    h.append(r)
    p._p.append(h)


def open_template():
    return Document(str(TEMPLATE))


# ------------------------------------------------------------------ pre-filled lists
def _pdf_pages(pdf):
    """[(page label, [lines])] for every page of a rendered PDF (poppler's pdftotext)."""
    import subprocess
    out = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True,
                         text=True, check=True).stdout
    pages = []
    for page in out.split("\f"):
        lines = [l.rstrip() for l in page.splitlines() if l.strip()]
        label = ""
        for l in reversed(lines):
            if "Thesis/Project" in l:
                label = l.split()[-1]
                break
        pages.append((label, lines))
    return pages


def _norm(s):
    return re.sub(r"\s+", " ", s).strip()


def _find_pages(texts, pages, *, start_at=0, caption=False):
    """Page label of each text, scanning forward in document order. Lines carrying dot
    leaders (the lists themselves) are never matched."""
    out, i = [], start_at
    for t in texts:
        key = _norm(t)
        hit = None
        for j in range(i, len(pages)):
            for l in pages[j][1]:
                l = _norm(l)
                if "....." in l or not l:
                    continue
                if caption:
                    ok = l.startswith(key.split(":")[0] + ":")
                else:
                    ok = l == key or (len(l) >= 18 and key.startswith(l))
                if ok:
                    hit = j
                    break
            if hit is not None:
                break
        if hit is None:
            out.append("")
        else:
            out.append(pages[hit][0]); i = hit
    return out


def _entry(text, page, style, indent):
    p = el("p")
    ppr_set(p, el("pStyle", val=style))
    tabs = el("tabs"); tabs.append(el("tab", val="right", leader="dot", pos=TEXT_W)); ppr_set(p, tabs)
    ppr_set(p, el("spacing", before=0, after=60, line=264, lineRule="auto"))
    ppr_set(p, el("ind", left=indent, right=600, hanging=0))
    p.append(make_run(text, bold=style == "TOC1"))
    p.append(make_run("\t" + page, bold=style == "TOC1"))
    return p


def fill_lists(doc, pdf):
    """Write real entries (with page numbers read from ``pdf``) into the TOC, List of
    Figures and List of Tables fields, so they show without a manual field update.
    Ctrl+A, F9 in Word still regenerates them."""
    body = doc.element.body
    pages = _pdf_pages(pdf)
    body_start = next((k for k, (_, ls) in enumerate(pages)
                       if any(_norm(l) == "Chapter 1" for l in ls)), 0)
    paras = [e for e in body if e.tag == W("p")]

    heads = []
    for e in paras:
        st = ppr(e).find(W("pStyle"))
        if st is None or st.get(W("val")) != "Heading1" or not ptext(e).strip():
            continue
        ol = ppr(e).find(W("outlineLvl"))
        heads.append((ptext(e).strip(), int(ol.get(W("val"))) + 1 if ol is not None else 1))
    hp = _find_pages([h for h, _ in heads], pages)
    caps = {"Figure": [], "Table": []}
    for e in paras:
        st = ppr(e).find(W("pStyle"))
        if st is not None and st.get(W("val")) == "Caption":
            t = _norm(ptext(e))
            caps[t.split()[0]].append(t)
    entries = {
        "TOC \\o": [( t, "TOC%d" % lvl, (lvl - 1) * 360, pg) for (t, lvl), pg in zip(heads, hp)],
        '"Figure"': [(t, "TableofFigures", 0, pg) for t, pg in
                     zip(caps["Figure"], _find_pages(caps["Figure"], pages, start_at=body_start, caption=True))],
        '"Table"': [(t, "TableofFigures", 0, pg) for t, pg in
                    zip(caps["Table"], _find_pages(caps["Table"], pages, start_at=body_start, caption=True))],
    }
    for e in paras:
        instr = "".join(t.text or "" for t in e.iter(W("instrText")))
        key = next((k for k in entries if k in instr), None)
        if key is None or not entries[key]:
            continue
        runs = [r for r in e if r.tag == W("r")]
        begin, ins, sep, _cached, end = runs[:5]
        items = [_entry(t, pg, st, ind) for t, st, ind, pg in entries[key]]
        first, last = items[0], items[-1]
        first.insert(1, sep); first.insert(1, ins); first.insert(1, begin)
        last.append(end)
        prev = e
        for it in items:
            prev.addnext(it); prev = it
        body.remove(e)
    return {k: (len(v), sum(1 for x in v if not x[3])) for k, v in entries.items()}
